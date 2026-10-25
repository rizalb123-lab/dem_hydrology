# -*- coding: utf-8 -*-
"""Rekapitulasi Excel (openpyxl)."""
from .i18n import tr
import datetime
import math

from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, ScatterChart, Reference, Series
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import registry as R
from . import report_common as C

HEAD = PatternFill("solid", fgColor="1F4E79")
SUB = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet_name(n):
    n = tr(n)
    for ch in '[]:*?/\\':
        n = n.replace(ch, '-')
    return n[:31]


def _hdr(ws, row, labels, widths=None):
    for j, t in enumerate(labels, 1):
        t = tr(t)
        c = ws.cell(row=row, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30


def _put(ws, row, vals, fmts=None, bold=False):
    for j, v in enumerate(vals, 1):
        if isinstance(v, float) and not math.isfinite(v):
            v = None
        c = ws.cell(row=row, column=j, value=v)
        c.border = BOX
        c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
        if bold:
            c.font = Font(bold=True)
        if fmts and j - 1 < len(fmts) and fmts[j - 1] and isinstance(v, (int, float)):
            c.number_format = fmts[j - 1]


def _title(ws, text, sub=None):
    text, sub = tr(text), tr(sub) if sub else sub
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


def build_excel(path, pl, d, project_title="Analisis Hidrologi DEM"):
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name("Ringkasan")
    _title(ws, project_title, tr("Dibuat %s oleh plugin Analisis Hidrologi DEM") % datetime.datetime.now().strftime("%d-%m-%Y %H:%M"))
    r = 4
    _hdr(ws, r, ["Keterangan", "Nilai"], [44, 90])
    for k, v in C.dem_summary(d, pl):
        r += 1
        _put(ws, r, [k, v])
    r += 2
    _hdr(ws, r, ["Analisis", "Status"])
    for k in pl.sel:
        r += 1
        _put(ws, r, [R.L()["ANALYSES"][k]["name"], pl.status.get(k, tr("tidak dijalankan"))])
    if pl.warnings:
        r += 2
        _hdr(ws, r, ["Peringatan", ""])
        for w in pl.warnings:
            r += 1
            _put(ws, r, [w, ""])
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.row_dimensions[r].height = 30
    r += 2
    _hdr(ws, r, ["Catatan penggunaan", ""])
    for t in C.method_note():
        r += 1
        _put(ws, r, [t, ""])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = 32

    # --- parameter
    ws = wb.create_sheet(_sheet_name("Parameter"))
    _title(ws, "Parameter yang dipakai")
    _hdr(ws, 3, ["Analisis", "Parameter", "Nilai", "Satuan"], [48, 52, 18, 12])
    for i, row in enumerate(C.params_rows(pl.par, pl.sel), 4):
        _put(ws, i, list(row))

    # --- statistik raster
    ws = wb.create_sheet(_sheet_name("Statistik"))
    _title(ws, "Statistik keluaran raster (di dalam batas studi)")
    cols = ["Keluaran", "Satuan", "Min", "P5", "P25", "Median", "P75", "P95", "Maks", "Rata-rata", "Simpangan baku", "Jumlah sel"]
    _hdr(ws, 3, cols, [46, 10] + [12] * 10)
    r = 3
    for o in pl.outs:
        if o.kind != "raster" or not o.stats:
            continue
        s = o.stats
        r += 1
        _put(ws, r, [o.title, o.unit, s["min"], s["p5"], s["p25"], s["median"], s["p75"], s["p95"], s["max"],
                     s["mean"], s["std"], s["n"]], [None, None] + ["0.000"] * 9 + ["#,##0"])
    ws.freeze_panes = "B4"

    # --- kelas
    ws = wb.create_sheet(_sheet_name("Kelas"))
    _title(ws, "Luas per kelas simbologi")
    _hdr(ws, 3, ["Keluaran", "Kelas", "Luas (ha)", "Persen (%)"], [46, 46, 14, 12])
    r = 3
    for o in pl.outs:
        for c in o.classes:
            r += 1
            _put(ws, r, [o.title, c["kelas"], c["luas_ha"], c["persen"]], [None, None, "#,##0.00", "0.0"])
    ws.freeze_panes = "A4"

    # --- sungai
    so = next((o for o in pl.outs if o.key == "stream_lines"), None)
    if so is not None:
        ws = wb.create_sheet(_sheet_name("Sungai"))
        _title(ws, "Jaringan sungai menurut orde Strahler (di dalam batas)")
        sm = so.extra["summary"]
        _hdr(ws, 3, ["Orde", "Jumlah ruas", "Panjang total (km)", "Panjang rata-rata (km)", "Rb (Nu/Nu+1)"], [10, 14, 20, 22, 16])
        for i, o_ in enumerate(so.extra["orders"], 4):
            _put(ws, i, [o_["orde"], o_["jumlah_ruas"], o_["panjang_total_km"], o_["panjang_rata2_km"], o_["Rb"]],
                 [None, None, "0.000", "0.000", "0.00"])
        n = len(so.extra["orders"])
        r = 4 + n + 1
        _put(ws, r, [tr("Ambang (ha)"), sm["ambang_ha"]], bold=True)
        _put(ws, r + 1, [tr("Total ruas"), sm["jumlah_ruas"]], bold=True)
        _put(ws, r + 2, [tr("Panjang total (km)"), sm["panjang_total_km"]], ["", "0.000"], bold=True)
        _put(ws, r + 3, [tr("Dd (km/km2)"), sm["Dd_km_km2"]], ["", "0.000"], bold=True)
        if n >= 2:
            ch = BarChart()
            ch.title = tr("Jumlah ruas per orde")
            ch.add_data(Reference(ws, min_col=2, min_row=3, max_row=3 + n), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n))
            ch.height, ch.width = 7.5, 14
            ws.add_chart(ch, "H3")

    # --- morfometri
    mo = next((o for o in pl.outs if o.key == "morphometry"), None)
    if mo is not None and mo.extra["rows"]:
        rows = mo.extra["rows"]
        ws = wb.create_sheet(_sheet_name("Morfometri"))
        _title(ws, "Morfometri DAS", "Setiap baris = satu unit. Rumus dan dasar tercantum pada baris keterangan di bawah tabel.")
        MM = C.morph_meta()
        heads = ["Unit", "Jenis"] + ["%s (%s)" % (m[1], m[2]) if m[2] != "-" else m[1] for m in MM] + \
                ["Bentuk DAS (Re)", "Tahap lahan (HI)", "Menyentuh tepi DEM"]
        _hdr(ws, 4, heads, [30, 18] + [16] * (len(heads) - 2))
        ws.row_dimensions[4].height = 58
        for i, m in enumerate(rows, 5):
            vals = [m["nama"], tr(m["jenis"])] + [m.get(k) for k, *_ in MM] + \
                   [C.interp_re(m.get("Re_elongasi")), C.interp_hi(m.get("HI")), tr("ya") if m.get("menyentuh_tepi_DEM") else tr("tidak")]
            _put(ws, i, vals, [None, None] + ["0.000"] * len(MM))
        ws.freeze_panes = "C5"
        r = 5 + len(rows) + 1
        ws.cell(row=r, column=1, value=tr("Keterangan rumus")).font = Font(bold=True)
        for k, lab, u, f, ref in MM:
            r += 1
            ws.cell(row=r, column=1, value=lab)
            ws.cell(row=r, column=2, value=f)
            ws.cell(row=r, column=3, value=R.REFS.get(ref, "")[:160])
        det = mo.extra["detail"]
        # --- orde per unit
        w2 = wb.create_sheet(_sheet_name("Orde per DAS"))
        _title(w2, "Jumlah dan panjang ruas per orde, per unit")
        _hdr(w2, 3, ["Unit", "Orde", "Jumlah ruas", "Panjang (km)"], [34, 8, 14, 14])
        r = 3
        for nm, dd in det.items():
            for o_, v in dd["orders"].items():
                r += 1
                _put(w2, r, [nm, o_, v["n"], v["len"] / 1000.0], [None, None, None, "0.000"])
        # --- hipsometri
        w3 = wb.create_sheet(_sheet_name("Hipsometri"))
        _title(w3, "Kurva hipsometrik (h/H terhadap a/A)")
        names = list(det.keys())
        col = 1
        first = None
        for nm in names[:8]:
            hy = det[nm]["hypso"]
            if not hy:
                continue
            w3.cell(row=3, column=col, value=nm).font = Font(bold=True)
            w3.cell(row=4, column=col, value="a/A")
            w3.cell(row=4, column=col + 1, value="h/H")
            for i, p in enumerate(hy, 5):
                w3.cell(row=i, column=col, value=p["a_rel"])
                w3.cell(row=i, column=col + 1, value=p["h_rel"])
            if first is None:
                first = (col, len(hy))
            col += 3
        if first:
            ch = ScatterChart()
            ch.title = tr("Kurva hipsometrik")
            ch.x_axis.title, ch.y_axis.title = "a/A", "h/H"
            ch.style = 13
            col2 = 1
            for nm in names[:8]:
                hy = det[nm]["hypso"]
                if not hy:
                    continue
                xs = Reference(w3, min_col=col2, min_row=5, max_row=4 + len(hy))
                ys = Reference(w3, min_col=col2 + 1, min_row=5, max_row=4 + len(hy))
                s = Series(ys, xs, title=nm)
                s.marker.symbol = "none"
                ch.series.append(s)
                col2 += 3
            ch.height, ch.width = 10, 16
            w3.add_chart(ch, "A30")
        # --- profil
        w4 = wb.create_sheet(_sheet_name("Profil Alur"))
        _title(w4, "Profil memanjang alur terpanjang (DEM terisi)")
        c0 = 1
        for nm in names[:6]:
            pf = det[nm]["profile"]
            if not pf:
                continue
            w4.cell(row=3, column=c0, value=nm).font = Font(bold=True)
            w4.cell(row=4, column=c0, value=tr("Jarak dari hulu (m)"))
            w4.cell(row=4, column=c0 + 1, value=tr("Elevasi (m)"))
            step = max(1, len(pf["dist"]) // 400)
            for i, (a, b) in enumerate(zip(pf["dist"][::step], pf["z"][::step]), 5):
                w4.cell(row=i, column=c0, value=a)
                w4.cell(row=i, column=c0 + 1, value=b)
            c0 += 3

    # --- time-area
    to = next((o for o in pl.outs if o.key == "travel_time"), None)
    if to is not None and to.extra.get("time_area"):
        ws = wb.create_sheet(_sheet_name("Time-Area"))
        _title(ws, "Fungsi luas-waktu (isochrone)", "Waktu tempuh ke titik keluar; untuk DAS titik outlet waktu dihitung ke outlet DAS tersebut.")
        col = 1
        for nm, rows in to.extra["time_area"].items():
            if not rows:
                continue
            ws.cell(row=3, column=col, value=nm).font = Font(bold=True)
            for j, t in enumerate([tr(x) for x in ["Waktu dari (menit)", "Waktu sampai (menit)", "Luas (ha)", "Kumulatif (ha)", "Kumulatif (%)"]]):
                c = ws.cell(row=4, column=col + j, value=t)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = HEAD
                ws.column_dimensions[get_column_letter(col + j)].width = 16
            for i, p in enumerate(rows, 5):
                for j, k in enumerate(["t0", "t1", "area_ha", "cum_ha", "cum_pct"]):
                    ws.cell(row=i, column=col + j, value=p[k]).number_format = "0.00"
            if col == 1:
                ch = LineChart()
                ch.title = tr("Kumulatif luas terhadap waktu (%s)") % nm
                ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=4 + len(rows)), titles_from_data=True)
                ch.set_categories(Reference(ws, min_col=2, min_row=5, max_row=4 + len(rows)))
                ch.height, ch.width = 8, 16
                ws.add_chart(ch, "H3")
            col += 6

    # --- tabel lain
    def simple(title, rows, sheet):
        if not rows:
            return
        w = wb.create_sheet(_sheet_name(sheet))
        _title(w, title)
        keys = list(rows[0].keys())
        _hdr(w, 3, keys, [18] * len(keys))
        for i, rw in enumerate(rows, 4):
            _put(w, i, [rw[k] for k in keys], ["0.000"] * len(keys))

    for key, title, sheet in [("subbasins", "Sub-DAS", "Sub-DAS"), ("watersheds", "DAS dari titik outlet", "DAS Titik"),
                              ("depressions", "Cekungan terbesar (maks. 50 baris)", "Cekungan")]:
        o = next((x for x in pl.outs if x.key == key), None)
        if o is not None:
            simple(title, o.extra.get("rows", []), sheet)
    ho = next((x for x in pl.outs if x.key == "hand_flood"), None)
    if ho is not None:
        simple("Skenario genangan HAND", ho.extra.get("stage_table", []), "Genangan HAND")
    bo = next((x for x in pl.outs if x.key == "bathtub"), None)
    if bo is not None and bo.extra.get("summary"):
        simple("Genangan bathtub", [bo.extra["summary"]], "Bathtub")
    # bobot indeks komposit
    wl = []
    for key in ("flood_index", "recharge_index", "erosion_index"):
        o = next((x for x in pl.outs if x.key == key), None)
        if o is not None:
            for f, w_ in o.extra.get("weights", {}).items():
                wl.append(dict(indeks=o.title, faktor=tr(f), bobot=w_))
    simple("Bobot indeks komposit", wl, "Bobot Indeks")

    # --- metode & pustaka
    ws = wb.create_sheet(_sheet_name("Metode & Pustaka"))
    _title(ws, "Dasar ilmiah setiap analisis")
    _hdr(ws, 3, ["Analisis", "Deskripsi", "Rumus / prosedur", "Batasan", "Pustaka"], [34, 56, 56, 56, 90])
    for i, k in enumerate(pl.sel, 4):
        a = R.L()["ANALYSES"][k]
        refs = "\n".join(t for _, t in R.refs_for(a["refs"]))
        _put(ws, i, [a["name"], a["desc"], a["formula"], a["limits"], refs])
        for j in range(1, 6):
            ws.cell(row=i, column=j).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = max(60, 15 * (refs.count("\n") + 2))
    wb.save(path)
    return path
